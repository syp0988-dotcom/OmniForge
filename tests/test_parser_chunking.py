"""Tests for the RAG document parser and structure-aware chunking."""


from agentflow.knowledge.chunking import (
    _select_strategy,
    chunk_by_code,
    chunk_by_markdown,
    chunk_by_paragraph,
    chunk_by_table,
    chunk_document,
    register_strategy,
)
from agentflow.knowledge.parser import (
    _read_zip,
    _read_raw_from_bytes,
    _strip_frontmatter,
    parse_document,
)


# -- raw parsing (bytes -> text) ---------------------------------------------


def test_read_txt_bytes():
    text = _read_raw_from_bytes("hello 中文".encode("utf-8"), "txt")
    assert "hello" in text and "中文" in text


def test_read_markdown_strips_frontmatter(tmp_path):
    raw = "---\ntitle: x\n---\n# 标题\n正文".encode("utf-8")
    # _read_raw_from_bytes only decodes text types; frontmatter stripping
    # happens in the file-based reader (parse_document for markdown).
    p = tmp_path / "probe_frontmatter.md"
    p.write_bytes(raw)
    text = parse_document(p, file_type="md", chunk_size=500, chunk_overlap=0)
    assert any("标题" in c for c in text)
    assert not any("title: x" in c for c in text)


def test_strip_frontmatter():
    assert _strip_frontmatter("---\na: 1\n---\nbody") == "body"
    assert _strip_frontmatter("no frontmatter") == "no frontmatter"


def test_read_html_strips_tags(tmp_path):
    p = tmp_path / "probe_page.html"
    p.write_bytes(b"<html><body><p>hello</p></body></html>")
    text = parse_document(p, file_type="html", chunk_size=500, chunk_overlap=0)
    assert any("hello" in c for c in text)
    assert not any("<p>" in c for c in text)


def test_read_csv_bytes():
    text = _read_raw_from_bytes(b"name,age\nAlice,30\nBob,25", "csv")
    assert "Alice" in text and "30" in text


# -- parse_document (file -> chunks) -----------------------------------------


def test_parse_document_txt(tmp_path):
    p = tmp_path / "note.txt"
    p.write_text("第一段。\n\n第二段。\n\n第三段。", encoding="utf-8")
    chunks = parse_document(p, chunk_size=50, chunk_overlap=0)
    assert len(chunks) >= 1
    assert all(isinstance(c, str) and c for c in chunks)


def test_parse_document_auto_detects_type(tmp_path):
    p = tmp_path / "doc.md"
    p.write_text("# 章节\n内容", encoding="utf-8")
    chunks = parse_document(p, chunk_size=100, chunk_overlap=10)
    assert chunks


# -- chunking strategies -----------------------------------------------------


def test_chunk_basic():
    parts = chunk_by_paragraph("word " * 120, chunk_size=50, overlap=10)
    assert len(parts) >= 2
    assert all(len(p) <= 55 for p in parts)


def test_chunk_by_markdown_preserves_headings():
    md = "# 第一章\n内容一\n\n# 第二章\n内容二\n"
    parts = chunk_by_markdown(md, chunk_size=500, overlap=0)
    assert any("第一章" in p for p in parts)
    assert any("第二章" in p for p in parts)


def test_chunk_by_code_splits_functions():
    code = "def a():\n    return 1\n\n\ndef b():\n    return 2\n"
    parts = chunk_by_code(code, chunk_size=500, overlap=0)
    joined = "\n".join(parts)
    assert "def a" in joined and "def b" in joined


def test_chunk_by_table():
    table = "| a | b |\n|---|---|\n| 1 | 2 |\n| 3 | 4 |\n"
    parts = chunk_by_table(table, chunk_size=500, overlap=0)
    assert any("1" in p and "2" in p for p in parts)


def test_select_strategy_mapping():
    assert _select_strategy("md") is chunk_by_markdown
    assert _select_strategy(".MD") is chunk_by_markdown
    assert _select_strategy("csv") is chunk_by_table
    assert _select_strategy("unknown") is not None


def test_register_strategy_extension():
    def custom(text, chunk_size=500, overlap=50):
        return ["custom:" + text]

    register_strategy("myext", custom)
    assert _select_strategy("myext") is custom
    assert chunk_document("x", "myext") == ["custom:x"]


def test_chunk_document_falls_back_for_unknown():
    parts = chunk_document("一段很长的文本。" * 20, "weird", chunk_size=50, overlap=5)
    assert parts


# -- binary formats (generated fixtures) -------------------------------------


def test_read_docx_binary(tmp_path):
    from docx import Document

    p = tmp_path / "report.docx"
    doc = Document()
    doc.add_heading("标题", level=1)
    doc.add_paragraph("这是正文内容")
    doc.save(str(p))

    text = parse_document(p, chunk_size=500, chunk_overlap=0)
    assert any("标题" in c for c in text)
    assert any("正文内容" in c for c in text)


def test_read_xlsx_binary(tmp_path):
    from openpyxl import Workbook

    p = tmp_path / "data.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "姓名"
    ws["A2"] = "张三"
    wb.save(str(p))

    text = parse_document(p, chunk_size=500, chunk_overlap=0)
    assert any("姓名" in c and "张三" in c for c in text)


def test_read_pptx_binary(tmp_path):
    from pptx import Presentation

    p = tmp_path / "deck.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "幻灯片标题"
    slide.placeholders[1].text = "要点内容"
    prs.save(str(p))

    text = parse_document(p, chunk_size=500, chunk_overlap=0)
    assert any("幻灯片标题" in c for c in text)


def test_read_pdf_binary_does_not_crash(tmp_path):
    from pypdf import PdfWriter

    p = tmp_path / "blank.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    with open(p, "wb") as fh:
        writer.write(fh)

    chunks = parse_document(p, chunk_size=500, chunk_overlap=0)
    assert isinstance(chunks, list)


def test_read_epub_binary(tmp_path):
    from ebooklib import epub

    p = tmp_path / "book.epub"
    book = epub.EpubBook()
    book.set_identifier("id123")
    book.set_title("测试书")
    book.set_language("zh")
    chapter = epub.EpubHtml(title="章", file_name="chap.xhtml", lang="zh")
    chapter.content = "<h1>第一章</h1><p>正文</p>"
    book.add_item(chapter)
    book.toc = (epub.Link("chap.xhtml", "章", "c1"),)
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    epub.write_epub(str(p), book)

    text = parse_document(p, chunk_size=500, chunk_overlap=0)
    assert any("第一章" in c for c in text)


def test_read_zip_returns_entries(tmp_path):
    import zipfile

    p = tmp_path / "bundle.zip"
    with zipfile.ZipFile(p, "w") as zf:
        zf.writestr("a.txt", "zip content")
    entries = _read_zip(p)
    assert isinstance(entries, list)
    assert any(name.endswith("a.txt") for name, _ in entries)
